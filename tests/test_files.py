import io
import json
from pathlib import Path
import zipfile

import httpx
import pytest

from test_client import SITE, TOKEN


def material(filename="lecture.pdf", url=None):
    return {"material_id": "a" * 20, "course_id": 569, "module_id": 32634,
            "filename": filename, "url": url or SITE + "/webservice/pluginfile.php/1/mod_resource/content/1/lecture.pdf",
            "modified_at": "2026-10-23T12:00:00+03:00", "size_bytes": None, "external": False}


def setup_store(tmp_path, handler):
    from university_mcp.client import MoodleClient
    from university_mcp.files import MaterialStore
    client = MoodleClient(TOKEN, tmp_path / "state", http=httpx.Client(transport=httpx.MockTransport(handler)))
    return MaterialStore(client, tmp_path / "downloads")


def tiny_pdf():
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({
        NameObject("/F1"): DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                            NameObject("/Subtype"): NameObject("/Type1"),
                                            NameObject("/BaseFont"): NameObject("/Helvetica")})})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 200 Td (Assets equal liabilities plus equity.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    result = io.BytesIO()
    writer.write(result)
    return result.getvalue()


def test_download_reads_real_pdf_with_provenance_and_bounded_pagination(tmp_path):
    def handle(request):
        assert request.url.params["token"] == TOKEN
        return httpx.Response(200, content=tiny_pdf(), headers={"Content-Type": "application/pdf"})
    store = setup_store(tmp_path, handle)
    result = store.read(material(), limit=24)
    assert result["engine"] == "markitdown" and result["text_format"] == "markdown"
    assert len(result["text"]) == 24 and result["next_offset"] == 24
    rest = store.read(material(), offset=24)
    assert "Assets equal liabilities plus equity." in result["text"] + rest["text"]
    assert rest["next_offset"] is None and rest["page_count"] == 1
    assert Path(result["local_path"]).read_bytes().startswith(b"%PDF-")
    assert TOKEN not in json.dumps(result)


@pytest.mark.parametrize("url", ["https://evil.example/file.pdf", "http://eng.sbcde.by/pluginfile.php/1/a.pdf",
                                  SITE + "/login/logout.php", SITE + "/mod/quiz/attempt.php?id=4"])
def test_download_refuses_unapproved_origins_and_non_file_routes(tmp_path, url):
    from university_mcp.client import PortalError
    def fail(request):
        pytest.fail("Unapproved target must be rejected before any request")
    store = setup_store(tmp_path, fail)
    with pytest.raises(PortalError):
        store.download(material(url=url))


def test_external_redirect_never_receives_token(tmp_path):
    from university_mcp.client import PortalError
    requests = []
    def handle(request):
        requests.append(request)
        return httpx.Response(302, headers={"Location": "https://external.example/download.pdf"})
    store = setup_store(tmp_path, handle)
    with pytest.raises(PortalError):
        store.download(material())
    assert len(requests) == 1
    assert list((tmp_path / "downloads").rglob("*.pdf")) == []


def test_login_html_is_not_accepted_as_a_pdf(tmp_path):
    from university_mcp.client import AuthenticationError
    store = setup_store(tmp_path, lambda r: httpx.Response(200, headers={"Content-Type": "text/html"},
                            text='<html><title>Log in to the site</title><form id="login"><input type="password"></form></html>'))
    with pytest.raises(AuthenticationError):
        store.download(material())
    assert list((tmp_path / "downloads").rglob("*.pdf")) == []


def test_filename_cannot_escape_download_directory(tmp_path):
    store = setup_store(tmp_path, lambda r: httpx.Response(200, content=tiny_pdf(), headers={"Content-Type": "application/pdf"}))
    result = store.download(material(filename="../../outside.pdf"))
    assert Path(result["local_path"]).is_relative_to(tmp_path / "downloads")
    assert not (tmp_path / "outside.pdf").exists()


def test_streaming_size_limit_removes_incomplete_file(tmp_path):
    from university_mcp.client import PortalError
    store = setup_store(tmp_path, lambda r: httpx.Response(200, content=b"x" * 100, headers={"Content-Type": "text/plain"}))
    store.max_download_bytes = 50
    with pytest.raises(PortalError, match="(?i)(size|large|limit)"):
        store.download(material(filename="big.txt"))
    assert list((tmp_path / "downloads").rglob("*.part")) == []


def test_docx_extracts_paragraphs_and_tables_without_images(tmp_path):
    from university_mcp.files import extract_text
    from docx import Document
    doc = Document()
    doc.add_paragraph("Homework requirements")
    doc.add_table(rows=1, cols=2).rows[0].cells[1].text = "Due Friday"
    path = tmp_path / "task.docx"
    doc.save(path)
    result = extract_text(path)
    assert "Homework requirements" in result["text"] and "Due Friday" in result["text"]
    assert result["extraction_status"] == "text"


def test_pptx_and_xlsx_extract_slide_order_and_sheet_values(tmp_path):
    from university_mcp.files import extract_text
    from pptx import Presentation
    from openpyxl import Workbook
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Banking lecture"
    ppt = tmp_path / "lecture.pptx"
    deck.save(ppt)
    assert "Banking lecture" in extract_text(ppt)["text"]
    book = Workbook()
    book.active.append(["Revenue", 420])
    xlsx = tmp_path / "case.xlsx"
    book.save(xlsx)
    assert "Revenue" in extract_text(xlsx)["text"] and "420" in extract_text(xlsx)["text"]


def test_unsupported_audio_returns_download_path_without_inventing_text(tmp_path):
    store = setup_store(tmp_path, lambda r: httpx.Response(200, content=b"ID3audio", headers={"Content-Type": "audio/mpeg"}))
    result = store.read(material(filename="recording.mp3"))
    assert result["extraction_status"] == "download_only" and result["text"] == ""
    assert Path(result["local_path"]).exists()


def test_archive_expansion_limit_is_checked_before_parsing_office_xml(tmp_path):
    from university_mcp.files import extract_text
    from university_mcp.client import PortalError
    path = tmp_path / "large.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "x" * 1000)
    with pytest.raises(PortalError, match="(?i)(archive|unpack|large|limit)"):
        extract_text(path, max_unpacked_bytes=100)


def test_markitdown_preserves_table_structure_for_agent_reading(tmp_path):
    from university_mcp.files import extract_text
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Metric", "Amount"
    table.cell(1, 0).text, table.cell(1, 1).text = "Revenue", "420"
    path = tmp_path / "table.docx"
    doc.save(path)
    result = extract_text(path)
    assert result["engine"] == "markitdown"
    assert "| Metric | Amount |" in result["text"]
    assert "| Revenue | 420 |" in result["text"]


def test_document_conversion_runs_once_and_subsequent_chunks_use_cache(tmp_path):
    from unittest.mock import patch
    import university_mcp.files as files
    calls = []
    def handle(request):
        calls.append(request)
        return httpx.Response(200, content=tiny_pdf(), headers={"Content-Type": "application/pdf"})
    store = setup_store(tmp_path, handle)
    original = files.extract_text
    with patch.object(files, "extract_text", wraps=original) as conversion:
        first = store.read(material(), limit=12)
        second = store.read(material(), offset=12, limit=12)
        assert first["engine"] == second["engine"] == "markitdown"
        assert first["markdown_path"] == second["markdown_path"]
        assert Path(first["markdown_path"]).read_text().startswith(first["text"] + second["text"])
        assert conversion.call_count == 1
    assert len(calls) == 1


@pytest.mark.parametrize("url", [
    SITE + "/webservice/pluginfile.php/../../login/logout.php",
    SITE + "/webservice/pluginfile.php/%2e%2e/%2e%2e/login/logout.php",
    SITE + "/webservice/pluginfile.php/%252e%252e/%252e%252e/login/logout.php",
])
def test_dot_segment_target_is_rejected_before_request_or_token_attachment(tmp_path, url):
    from university_mcp.client import PortalError
    def fail(request):
        pytest.fail("Traversal route was requested: " + request.url.path)
    store = setup_store(tmp_path, fail)
    with pytest.raises(PortalError):
        store.download(material(url=url))


def test_cached_markdown_is_updated_when_source_revision_changes(tmp_path):
    content = [b"old lecture text"]
    store = setup_store(tmp_path, lambda r: httpx.Response(200, content=content[0], headers={"Content-Type": "text/plain"}))
    meta = material(filename="lecture.txt")
    first = store.read(meta)
    assert Path(first["markdown_path"]).read_text() == "old lecture text"
    content[0] = b"new lecture text"
    updated = {**meta, "modified_at": "2026-10-24T12:00:00+03:00"}
    second = store.read(updated)
    assert second["text"] == "new lecture text"
    assert Path(second["markdown_path"]).read_text() == "new lecture text"


def test_long_unicode_filename_allows_markdown_and_metadata_sidecars(tmp_path):
    store = setup_store(tmp_path, lambda r: httpx.Response(200, content=b"Lecture text", headers={"Content-Type": "text/plain"}))
    result = store.read(material(filename="Длиннаялекция" * 20 + ".txt"))
    assert Path(result["markdown_path"]).read_text() == "Lecture text"
    assert len(Path(result["markdown_path"]).name.encode("utf-8")) < 256

def test_cached_unicode_text_uses_utf8_even_with_legacy_system_encoding(tmp_path,monkeypatch):
    """Windows cp1252 must not determine the encoding of our JSON/Markdown files."""
    import tempfile
    from university_mcp.files import _atomic_text
    original=tempfile.NamedTemporaryFile
    def legacy_default(*args,**kwargs):
        kwargs.setdefault('encoding','cp1252')
        return original(*args,**kwargs)
    monkeypatch.setattr(tempfile,'NamedTemporaryFile',legacy_default)
    path=tmp_path/'material.md'
    _atomic_text(path,'Лекция об учёте')
    assert path.read_bytes()=='Лекция об учёте'.encode('utf-8')
